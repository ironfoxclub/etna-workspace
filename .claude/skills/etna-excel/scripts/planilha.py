"""Gera e lê planilhas Excel (.xlsx).
Uso:
  uv run --with openpyxl planilha.py ler arquivo.xlsx
  uv run --with openpyxl planilha.py exemplo saida.xlsx
  uv run --with openpyxl planilha.py criar saida.xlsx dados.json
dados.json: {"abas": {"Vendas": [["Produto","Qtd","Preço"],["A",2,10.5]]}}
Fórmulas entram como texto começando com "=" (ex.: "=B2*C2").
"""
import json, sys
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill


def criar(destino, dados):
    wb = Workbook(); wb.remove(wb.active)
    for nome, linhas in dados["abas"].items():
        ws = wb.create_sheet(nome[:31])
        for linha in linhas: ws.append(linha)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="333333")
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(
                50, max(len(str(c.value or "")) for c in col) + 2)
    wb.save(destino)


def ler(caminho):
    # mostra a fórmula e, se o Excel já calculou e salvou, o valor (fórmula→valor)
    val = load_workbook(caminho, data_only=True)
    frm = load_workbook(caminho)
    for ws in frm:
        print(f"--- {ws.title} ({ws.max_row}x{ws.max_column}) ---")
        for linha in ws.iter_rows():
            partes = []
            for c in linha:
                v = val[ws.title][c.coordinate].value
                partes.append(f"{c.value}" if v is None or v == c.value else f"{c.value}→{v}")
            print(" | ".join(partes))


EXEMPLO = {"abas": {"Vendas": [["Produto", "Qtd", "Preço", "Total"],
                               ["Camiseta", 3, 40, "=B2*C2"], ["Boné", 2, 25, "=B3*C3"],
                               ["Soma", "", "", "=SUM(D2:D3)"]]}}

if __name__ == "__main__":
    if len(sys.argv) < 3: sys.exit(__doc__)
    cmd, arq = sys.argv[1], sys.argv[2]
    if cmd == "ler": ler(arq)
    elif cmd == "exemplo": criar(arq, EXEMPLO)
    elif cmd == "criar": criar(arq, json.load(open(sys.argv[3], encoding="utf-8")))
    else: sys.exit(__doc__)
