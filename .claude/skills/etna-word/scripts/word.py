"""Gera e lê arquivos Word (.docx).
Uso:
  uv run --with python-docx word.py ler arquivo.docx
  uv run --with python-docx word.py exemplo saida.docx
  uv run --with python-docx word.py criar saida.docx conteudo.json
conteudo.json: {"titulo": "...", "blocos": [{"h1": "..."}, {"p": "..."}, {"lista": ["a","b"]},
                {"tabela": [["Col A","Col B"],["1","2"]]}]}
"""
import json, sys
from docx import Document


def criar(destino, dados):
    doc = Document()
    doc.add_heading(dados.get("titulo", "Documento"), 0)
    for b in dados.get("blocos", []):
        if "h1" in b: doc.add_heading(b["h1"], 1)
        elif "h2" in b: doc.add_heading(b["h2"], 2)
        elif "p" in b: doc.add_paragraph(b["p"])
        elif "lista" in b:
            for item in b["lista"]: doc.add_paragraph(item, style="List Bullet")
        elif "tabela" in b:
            linhas = b["tabela"]
            t = doc.add_table(rows=len(linhas), cols=len(linhas[0]), style="Table Grid")
            for i, linha in enumerate(linhas):
                for j, cel in enumerate(linha): t.cell(i, j).text = str(cel)
    doc.save(destino)


def ler(caminho):
    doc = Document(caminho)
    for p in doc.paragraphs:
        if p.text.strip(): print(f"[{p.style.name}] {p.text}")
    for n, t in enumerate(doc.tables, 1):
        print(f"--- tabela {n} ---")
        for linha in t.rows: print(" | ".join(c.text for c in linha.cells))


EXEMPLO = {"titulo": "Proposta de exemplo", "blocos": [
    {"h1": "Resumo"}, {"p": "Texto de abertura da proposta."},
    {"lista": ["Primeiro ponto", "Segundo ponto"]},
    {"tabela": [["Item", "Valor"], ["Site", "1000"], ["Anúncios", "500"]]}]}

if __name__ == "__main__":
    if len(sys.argv) < 3: sys.exit(__doc__)
    cmd, arq = sys.argv[1], sys.argv[2]
    if cmd == "ler": ler(arq)
    elif cmd == "exemplo": criar(arq, EXEMPLO)
    elif cmd == "criar": criar(arq, json.load(open(sys.argv[3], encoding="utf-8")))
    else: sys.exit(__doc__)
