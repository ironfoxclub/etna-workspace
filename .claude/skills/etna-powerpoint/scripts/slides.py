"""Gera e lê apresentações PowerPoint (.pptx).
Uso:
  uv run --with python-pptx slides.py ler arquivo.pptx
  uv run --with python-pptx slides.py exemplo saida.pptx
  uv run --with python-pptx slides.py criar saida.pptx slides.json
slides.json: {"slides": [{"titulo": "...", "subtitulo": "..."},
                         {"titulo": "...", "topicos": ["a","b"], "notas": "..."},
                         {"titulo": "...", "imagem": "foto.png"}]}
O primeiro slide sem tópicos nem imagem vira capa.
"""
import json, sys
from pptx import Presentation
from pptx.util import Inches


def criar(destino, dados):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)  # 16:9
    for i, s in enumerate(dados["slides"]):
        if i == 0 and "topicos" not in s and "imagem" not in s:
            sl = prs.slides.add_slide(prs.slide_layouts[0])
            sl.shapes.title.text = s["titulo"]
            sl.placeholders[1].text = s.get("subtitulo", "")
        else:
            sl = prs.slides.add_slide(prs.slide_layouts[5 if "imagem" in s else 1])
            sl.shapes.title.text = s["titulo"]
            if "topicos" in s:
                tf = sl.placeholders[1].text_frame
                tf.text = s["topicos"][0]
                for t in s["topicos"][1:]: tf.add_paragraph().text = t
            if "imagem" in s:
                sl.shapes.add_picture(s["imagem"], Inches(2), Inches(1.8), height=Inches(5))
        if s.get("notas"): sl.notes_slide.notes_text_frame.text = s["notas"]
    prs.save(destino)


def ler(caminho):
    prs = Presentation(caminho)
    for n, sl in enumerate(prs.slides, 1):
        print(f"--- slide {n} ---")
        for sh in sl.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip(): print(sh.text_frame.text)
        if sl.has_notes_slide and sl.notes_slide.notes_text_frame.text.strip():
            print("[notas]", sl.notes_slide.notes_text_frame.text)


EXEMPLO = {"slides": [{"titulo": "Plano do trimestre", "subtitulo": "Exemplo Etna"},
                      {"titulo": "Metas", "topicos": ["Dobrar leads", "Baixar o custo por lead"], "notas": "Falar devagar."}]}

if __name__ == "__main__":
    if len(sys.argv) < 3: sys.exit(__doc__)
    cmd, arq = sys.argv[1], sys.argv[2]
    if cmd == "ler": ler(arq)
    elif cmd == "exemplo": criar(arq, EXEMPLO)
    elif cmd == "criar": criar(arq, json.load(open(sys.argv[3], encoding="utf-8")))
    else: sys.exit(__doc__)
